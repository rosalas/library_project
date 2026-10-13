import json
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from logger import logger


class ReportGenerator:
    def __init__(self, output_dir: Path | None = None):
        self.output_dir = output_dir or Path("data/outputs")
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_excel_dashboard(
        self,
        catalog_df: pd.DataFrame,
        author_df: pd.DataFrame,
        filename: str = "Library_Catalog_Report.xlsx",
    ) -> Path:
        # Build a multi-tab formatted Excel workbook containing catalog data, author analytics, and KPI formulas
        filepath = self.output_dir / filename
        wb = Workbook()

        # Tab 1: Catalog overview
        ws_catalog = wb.active
        ws_catalog.title = "Catalog Overview"
        self._write_df_to_sheet(ws_catalog, catalog_df)

        # Tab 2: Author analytics
        ws_authors = wb.create_sheet(title="Author Analytics")
        self._write_df_to_sheet(ws_authors, author_df)

        # Tab 3: Summary metrics (KPI dashboard)
        ws_summary = wb.create_sheet(title="Executive Summary")
        self._build_executive_summary_tab(ws_summary, len(catalog_df), len(author_df))

        wb.save(filepath)
        logger.info("Successfully exported Excel dashboard to '%s'", filepath)
        return filepath

    def _write_df_to_sheet(self, ws: Worksheet, df: pd.DataFrame) -> None:
        # Write a pandas DataFrame to an openpyxl worksheet with custom headers, borders, gridlines, and auto-fit columns
        # Enable visible gridlines in Excel
        ws.views.sheetView[0].showGridLines = True

        # Style definitions
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        if df.empty:
            ws.append(["No data available"])
            return

        # 1. Append headers & apply formatting
        ws.append(list(df.columns))
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 2. Append data rows
        for row in df.itertuples(index=False):
            ws.append(list(row))

        # 3. Apply cell borders and auto-fit column widths
        for col_idx, col_cells in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in col_cells:
                if cell.row > 1:
                    cell.border = thin_border
                val_str = str(cell.value or "")
                max_len = max(max_len, len(val_str))

            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    def _build_executive_summary_tab(
        self, ws: Worksheet, catalog_row_count: int, author_row_count: int
    ) -> None:
        # Construct the KPI Executive Summary worksheet using native Excel formulas
        ws.views.sheetView[0].showGridLines = True

        # Header styles
        title_font = Font(name="Calibri", size=16, bold=True, color="1F4E78")
        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        border_thin = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        # Title block
        ws["A1"] = "Library Catalog Key Performance Indicators"
        ws["A1"].font = title_font

        # KPI table headers
        ws["A3"] = "Metric Indicator"
        ws["B3"] = "Calculated Value (Formula)"
        ws["A3"].fill = header_fill
        ws["A3"].font = header_font
        ws["B3"].fill = header_fill
        ws["B3"].font = header_font

        # Calculate exact target ranges for Excel formulas
        catalog_end_row = max(catalog_row_count + 1, 2)
        author_end_row = max(author_row_count + 1, 2)

        metrics = [
            (
                "Total Books Processed in Catalog",
                f"=COUNTA('Catalog Overview'!A2:A{catalog_end_row})",
            ),
            (
                "Total Unique Authors Cataloged",
                f"=COUNTA('Author Analytics'!A2:A{author_end_row})",
            ),
            (
                "Average First Publication Year",
                f"=AVERAGE('Catalog Overview'!C2:C{catalog_end_row})",
            ),
            (
                "Earliest Book Publication Year",
                f"=MIN('Catalog Overview'!C2:C{catalog_end_row})",
            ),
            (
                "Latest Book Publication Year",
                f"=MAX('Catalog Overview'!C2:C{catalog_end_row})",
            ),
        ]

        # Write KPI rows
        for idx, (label, formula) in enumerate(metrics, start=4):
            ws.cell(row=idx, column=1, value=label).border = border_thin
            val_cell = ws.cell(row=idx, column=2, value=formula)
            val_cell.border = border_thin
            val_cell.alignment = Alignment(horizontal="right")

        # Column sizing
        ws.column_dimensions["A"].width = 38
        ws.column_dimensions["B"].width = 28

    def export_csv(self, df: pd.DataFrame, filename: str) -> Path:
        # Export a pandas DataFrame as a UTF-8 encoded CSV file
        filepath = self.output_dir / filename
        df.to_csv(filepath, index=False, encoding="utf-8")
        logger.info("Successfully exported CSV to '%s'", filepath)
        return filepath

    def export_json_summary(
        self, analytics_dict: dict[str, Any], filename: str
    ) -> Path:
        # Export analytics dictionary wrapped with generation metadata as a formatted JSON file
        filepath = self.output_dir / filename

        payload = {
            "metadata": {
                "generated_at": datetime.now(UTC).isoformat(),
                "generator": "OLICAS ReportGenerator v1.0",
            },
            "summary_data": analytics_dict,
        }

        with open(filepath, mode="w", encoding="utf-8") as f:
            json.dump(payload, f, indent=2, ensure_ascii=False)

        logger.info("Successfully exported JSON summary to '%s'", filepath)
        return filepath
